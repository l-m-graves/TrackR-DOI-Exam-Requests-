"""Settings written the way the real local_settings.py writes them: a request-ID pattern that
allows 'FIN 45', the company's own domain in SENDER_DOMAINS, whole addresses where domains are
expected, MUST_BE_SENT_TO for external emails, and LOG_PATH given as a folder. Fake values only."""
import re
import shutil
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest import mock

import doi_filter
from attachment_text import read_text
from doi_analysis import detect
from doi_filter import Mail, canonical_id, log_file, normalize_domains
from tests import fake_doi_emails as E

LOOSE_PATTERN = r"\b(?:FIN|IT)[ -]?\d+\b|\bDOI-[A-Z0-9]+-\d+\b"     # the shape used in real settings


class LooseIdTests(unittest.TestCase):
    def setUp(self):
        p = mock.patch.object(doi_filter, "REQUEST_ID", re.compile(LOOSE_PATTERN))
        p.start()
        self.addCleanup(p.stop)

    def test_every_spelling_becomes_one_id(self):
        for text in ("FIN-45 Request", "RE: FIN 45 Request", "FIN45 support", "Request FIN-45:"):
            self.assertEqual(doi_filter.find_id(text), "FIN-45", text)
        self.assertEqual(doi_filter.find_id("IT 39 Access Reviews"), "IT-39")
        self.assertEqual(doi_filter.find_id("DOI-TEST-07 Request"), "DOI-TEST-07")

    def test_ordinary_words_are_not_ids(self):
        self.assertEqual(doi_filter.find_id("Please send it 5 days early"), "")     # case-sensitive
        self.assertEqual(doi_filter.find_id("Final 2027 numbers"), "")

    def test_tracker_text_gets_the_hyphenated_id_first(self):
        import upsert_tracker
        with mock.patch.object(upsert_tracker, "REQUEST_ID", re.compile(LOOSE_PATTERN)):
            self.assertEqual(upsert_tracker.request_text("RE: FIN 45 Request - Payroll", "FIN-45"),
                             "FIN-45 Request - Payroll")
            self.assertEqual(upsert_tracker.first_word_id("FIN-45 Request - Payroll"), "FIN-45")

    def test_canonical_id_leaves_hyphenated_ids_alone(self):
        self.assertEqual(canonical_id("FIN-45"), "FIN-45")
        self.assertEqual(canonical_id("fin 45"), "FIN-45")
        self.assertEqual(canonical_id("DOI-TEST-01"), "DOI-TEST-01")


class DomainSettingTests(unittest.TestCase):
    def test_whole_addresses_count_as_their_domain(self):
        self.assertEqual(normalize_domains({"Someone@Exam-Firm.test", "@doi-state.test", "corp.test", " "}),
                         {"exam-firm.test", "doi-state.test", "corp.test"})

    def test_own_company_in_sender_domains_still_counts_as_internal(self):
        with mock.patch.multiple(doi_filter, **dict(E.ROLES, SENDER_DOMAINS={"corp.test", "doi-state.test"})):
            self.assertEqual(doi_filter.role_of("colleague@corp.test"), "internal")
            self.assertEqual(doi_filter.role_of("ana.accountant@corp.test"), "accountant")
            self.assertEqual(doi_filter.role_of("someone@doi-state.test"), "examiner")
        with mock.patch.multiple(doi_filter, **dict(E.ROLES, SENDER_DOMAINS={"legacy-doi.test"}, EXAMINER_DOMAINS=set())):
            self.assertEqual(doi_filter.role_of("x@legacy-doi.test"), "examiner")     # legacy list still works


class MustBeSentToTests(unittest.TestCase):
    MAILBOX = "audit@corp.test"

    def setUp(self):
        p = mock.patch.multiple(doi_filter, **dict(E.ROLES, MUST_BE_SENT_TO={self.MAILBOX}))
        p.start()
        self.addCleanup(p.stop)

    def judge(self, sender, to, cc=(E.INTERN,)):
        return detect(E.mail("FIN-401 Request - Premium Tax",
                             "Please provide the DOI examination support for premium tax.", sender, to=to, cc=cc))

    def test_external_email_must_reach_the_mailbox(self):
        self.assertEqual(self.judge("lead.examiner@doi-state.test", [self.MAILBOX]).status, "accept")
        det = self.judge("lead.examiner@doi-state.test", ["ana.accountant@corp.test"])
        self.assertEqual(det.status, "reject")
        self.assertIn("MUST_BE_SENT_TO", det.reasons[0])

    def test_internal_forward_to_you_is_not_held_to_it(self):
        det = self.judge("ana.accountant@corp.test", [E.INTERN], cc=())
        self.assertEqual(det.status, "accept")


class LogFolderTests(unittest.TestCase):
    def test_folder_gets_the_standard_file_name(self):
        self.assertEqual(Path(log_file("/share/DOI Requests Tracker/Log")),
                         Path("/share/DOI Requests Tracker/Log/DOI_Run_Log.csv"))
        self.assertEqual(log_file("/share/Logs/My_Log.csv"), "/share/Logs/My_Log.csv")
        self.assertIsNone(log_file(None))


class CsvAttachmentTests(unittest.TestCase):
    def test_csv_attachments_are_read(self):
        d = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, d)
        path = d / "FIN-45_Request_Items.csv"
        path.write_text("Item,Description\n1,Premium tax returns for 2026\n", encoding="utf-8")
        text, status = read_text(path)
        self.assertEqual(status, "read")
        self.assertIn("Premium tax returns", text)


if __name__ == "__main__":
    unittest.main()


class OutlookInternalForwardTests(unittest.TestCase):
    """Classic Outlook quotes colleagues as 'From: Lastname, Firstname' - no address."""
    BODY = ("Adding the team - see below.\n\n"
            "From: Lee, Jordan\n"
            "Sent: Thursday, October 7, 2027 9:12 AM\n"
            "To: Accountant, Ana <ana.accountant@corp.test>\n"
            "Subject: RE: FIN-501 Request - Payroll\n\n"
            "We expect to provide the payroll support by October 20, 2027.\n")

    def setUp(self):
        p = mock.patch.multiple(doi_filter, **E.ROLES)
        p.start()
        self.addCleanup(p.stop)

    def analyze(self, names=None, body=BODY, attachments=None):
        from doi_analysis import analyze
        m = E.mail("RE: FIN-501 Request - Payroll", body, "ana.accountant@corp.test",
                   received=datetime(2027, 10, 7, 10), names=names or {})
        m.attachment_texts = attachments or {}
        return analyze(m, "FIN-501")

    def test_name_found_on_the_email_gives_the_address_and_role(self):
        an = self.analyze(names={"jordan.lee@corp.test": "Jordan Lee"})
        self.assertEqual(f"{an.internal_expected_date:%Y-%m-%d}", "2027-10-20")
        self.assertIsNone(an.exam_due_date)                    # a colleague's date is never the DOI's
        self.assertFalse(any("not in the role settings" in r for r in an.review))

    def test_unknown_name_is_labelled_as_the_forwarded_part_not_the_forwarder(self):
        an = self.analyze()
        note = next(r for r in an.review if "not in the role settings" in r)
        self.assertIn("forwarded part from Lee, Jordan", note)
        self.assertNotIn("ana.accountant", note)

    def test_exchange_address_means_internal(self):
        body = self.BODY.replace("From: Lee, Jordan", "From: Lee, Jordan </O=EXCHANGELABS/OU=EXCHANGE/CN=RECIPIENTS/CN=JLEE>")
        an = self.analyze(body=body)
        self.assertEqual(f"{an.internal_expected_date:%Y-%m-%d}", "2027-10-20")

    def test_attachment_named_with_the_request_id_is_a_request_document(self):
        doc = {"2027-10-07_Sample Group Exam - FIN-501 Payroll Support.docx":
               "FIN-501 Payroll. Please provide the payroll registers by October 22, 2027."}
        an = self.analyze(body="Please see the attached DOI request.", attachments=doc)
        self.assertEqual(f"{an.exam_due_date:%Y-%m-%d}", "2027-10-22")
        self.assertTrue(doi_filter.important_attachments(list(doc)))
        self.assertFalse(doi_filter.important_attachments(["Quarterly newsletter.pdf"]))


class ContactTests(unittest.TestCase):
    """Column L = the person inside the company asked to provide the information."""
    ME = "intern@corp.test"

    def setUp(self):
        p = mock.patch.multiple(doi_filter, **dict(E.ROLES, MUST_BE_SENT_TO={"audit@corp.test"},
                                                   MY_ADDRESSES={self.ME}))
        p.start()
        self.addCleanup(p.stop)

    def contact(self, sender, to, kind="other", names=None):
        from doi_analysis import analyze
        from upsert_tracker import contact
        m = E.mail("FIN-601 Request - Claims", "Please provide the DOI examination claims listing.", sender,
                   to=to, cc=(), names=names or {})
        return contact(m, kind, analyze(m, "FIN-601").asked)

    def test_the_person_asked_inside_the_company(self):
        self.assertEqual(self.contact("ana.accountant@corp.test", ["pat.claims@corp.test"],
                                      names={"pat.claims@corp.test": "Pat Claims"}),
                         "Pat Claims <pat.claims@corp.test>")

    def test_never_you_the_mailbox_an_accountant_or_the_examiner(self):
        self.assertEqual(self.contact("ana.accountant@corp.test", [self.ME]), "")         # forwarded to you
        self.assertEqual(self.contact("ana.accountant@corp.test",
                                      [self.ME, "audit@corp.test", "ana.accountant@corp.test",
                                       "lead.examiner@doi-state.test", "pat.claims@corp.test"]),
                         "pat.claims@corp.test")
        self.assertEqual(self.contact("lead.examiner@doi-state.test", ["audit@corp.test"]), "")   # examiner's email

    def test_a_cell_holding_only_your_address_may_be_replaced(self):
        from upsert_tracker import only_you
        self.assertTrue(only_you(f"Intern Name <{self.ME}>"))
        self.assertFalse(only_you(f"Intern Name <{self.ME}>; pat.claims@corp.test"))
        self.assertFalse(only_you("Pat Claims"))                                  # typed by a person

    def test_own_addresses_from_outlook_accounts(self):
        from types import SimpleNamespace
        ns = SimpleNamespace(Accounts=[SimpleNamespace(SmtpAddress="Intern@Corp.test")])   # no CurrentUser: fine
        self.assertEqual(doi_filter.own_addresses(ns), {"intern@corp.test"})


class ExaminerThreadTests(unittest.TestCase):
    """Outside people the examiner side puts on a request thread count as third party."""
    FIRM = "kim.reviewer@newfirm.test"          # exam-firm person whose domain isn't in the settings

    def setUp(self):
        p = mock.patch.multiple(doi_filter, **dict(E.ROLES, MUST_BE_SENT_TO={"audit@corp.test"},
                                                   MY_ADDRESSES={E.INTERN}, THREAD_THIRD_PARTIES=set()))
        p.start()
        self.addCleanup(p.stop)

    def included(self, sender, to=(), cc=(), body="Please provide the DOI examination items."):
        from doi_analysis import examiner_included
        return examiner_included(E.mail("FIN-701 Request - Claims", body, sender, to=to, cc=cc))

    def test_people_the_examiner_copies(self):
        got = self.included("lead.examiner@doi-state.test", to=["audit@corp.test", "ana.accountant@corp.test"],
                            cc=[self.FIRM, E.INTERN, "pat.claims@corp.test"])
        self.assertEqual(got, {self.FIRM})               # not you, the mailbox, colleagues

    def test_recipients_of_a_forwarded_examiner_message(self):
        body = ("FYI - see below.\n\n"
                "From: Lead Examiner <lead.examiner@doi-state.test>\n"
                "Sent: Monday, October 4, 2027 9:00 AM\n"
                "To: Audit <audit@corp.test>\n"
                f"Cc: Kim Reviewer <{self.FIRM}>\n"
                "Subject: FIN-701 Request - Claims\n\n"
                "Please provide the claims listing.\n")
        self.assertEqual(self.included("ana.accountant@corp.test", to=[E.INTERN], body=body), {self.FIRM})

    def test_people_a_colleague_copies_are_not_vouched_for(self):
        self.assertEqual(self.included("ana.accountant@corp.test", cc=["consultant@advisers.test"]), set())

    def test_role_and_contact_directory(self):
        import doi_contacts as C
        contacts = {}
        m = E.mail("FIN-701 Request - Claims", "Please provide the items.", "lead.examiner@doi-state.test",
                   to=["audit@corp.test"], cc=[self.FIRM])
        C.observe(contacts, m, "FIN-701", datetime(2027, 10, 4), {self.FIRM})
        self.assertEqual(contacts[self.FIRM]["Examiner_Thread"], "Y")
        self.assertEqual(C.thread_third_parties(contacts), {self.FIRM})
        with mock.patch.object(doi_filter, "THREAD_THIRD_PARTIES", C.thread_third_parties(contacts)):
            self.assertEqual(doi_filter.role_of(self.FIRM), "third_party")
            self.assertEqual(doi_filter.role_of("ana.accountant@corp.test"), "accountant")
        reply = E.mail("RE: FIN-701", "Attached.", self.FIRM, to=["ana.accountant@corp.test"])
        C.observe(contacts, reply, "FIN-701", datetime(2027, 10, 5))
        self.assertNotIn(self.FIRM, C.pending_senders(contacts))     # trusted via the examiner
        contacts[self.FIRM]["Approved"] = "N"                        # you say no
        self.assertEqual(C.thread_third_parties(contacts), set())
        self.assertNotIn(self.FIRM, C.pending_senders(contacts))     # decided - not asked again

    def test_forward_only_addresses_get_a_row(self):
        import doi_contacts as C
        contacts = {}
        m = E.mail("FW: FIN-701", "x", "ana.accountant@corp.test", to=[E.INTERN])
        C.observe(contacts, m, "FIN-701", datetime(2027, 10, 4), {self.FIRM})
        self.assertEqual((contacts[self.FIRM]["Roles"], contacts[self.FIRM]["Examiner_Thread"]), ("Forwarded", "Y"))


class ExaminerThreadRunTests(unittest.TestCase):
    """Two hourly runs: the examiner copies an unknown exam-firm person; that person's own
    email waits in review, and is accepted on the next run."""

    def test_review_becomes_accepted_once_the_examiner_vouched(self):
        import io
        from contextlib import redirect_stdout
        import run
        from tests.test_hourly import FakeFolder, fake_msg, fake_outlook, person
        d = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, d)
        firm = "kim.reviewer@newfirm.test"
        examiner = fake_msg("FIN-702 Request - Claims", "Please provide the DOI examination claims listing.",
                            "lead.examiner@doi-state.test", "E1", "T1")
        examiner.Recipients.append(person(firm, 2))
        reply = fake_msg("RE: FIN-702 Request - Claims", "Following up on the DOI examination request for claims.",
                         firm, "E2", "T1")
        inbox = FakeFolder("Inbox", [examiner, reply])
        uploads = []
        def fake_upsert(path, accepted, **kw):
            uploads.append(sorted(m.entry_id for m, _ in accepted))
            return [{"request_id": rid, "row": 4, "action": "updated", "ref": doi_filter.message_ref(m.entry_id),
                     "mail": m, "changes": "note"} for m, rid in accepted]
        with mock.patch.multiple(doi_filter, **dict(E.ROLES, MUST_BE_SENT_TO=set(), THREAD_THIRD_PARTIES=set(),
                                                    LOG_PATH=str(d / "DOI_Run_Log.csv"), TRACKER_PATH=str(d / "T.xlsx"),
                                                    CONTACTS_PATH=str(d / "DOI_Contacts.csv"), PUBLISH_PATH=None,
                                                    ATTACH_DIR=None, TRACK_SENT=False)), \
                mock.patch.object(doi_filter, "_outlook", return_value=fake_outlook(inbox)), \
                mock.patch.object(run, "upsert", side_effect=fake_upsert), redirect_stdout(io.StringIO()) as out:
            run.main([])
            run.main([])
        self.assertEqual(uploads, [["E1"], ["E2"]])          # run 1: examiner only; run 2: the firm's reply
        self.assertIn("now count as third party", out.getvalue())


class SendToTests(unittest.TestCase):
    """Who the response goes to: the examiner side's own words, else the lead examiner."""

    def setUp(self):
        p = mock.patch.multiple(doi_filter, **E.ROLES)
        p.start()
        self.addCleanup(p.stop)

    def send_to(self, body, sender="lead.examiner@doi-state.test", to=("audit@corp.test",),
                cc=("jordan.lee@examfirm.test",), names=None, leads=None):
        from doi_analysis import analyze
        m = E.mail("FIN-801 Request - Claims", body, sender, to=to, cc=cc,
                   names=names if names is not None else {"jordan.lee@examfirm.test": "Lee, Jordan"})
        with mock.patch.object(doi_filter, "LEAD_EXAMINERS",
                               leads if leads is not None else {"lead.examiner@doi-state.test"}):
            an = analyze(m, "FIN-801")
        return an.send_to, an.send_to_source

    def test_named_person_is_looked_up_on_the_email(self):
        who, src = self.send_to("Please send the responses to Jordan by 10/20/27.")
        self.assertEqual(who, "Lee, Jordan <jordan.lee@examfirm.test>")
        self.assertTrue(src.startswith("stated: "))

    def test_me_is_whoever_wrote_it(self):
        self.assertEqual(self.send_to("Please give it to me by 10/20/27.")[0], "lead.examiner@doi-state.test")

    def test_me_in_a_forwarded_examiner_message(self):
        body = ("FYI below.\n\nFrom: Lead Examiner <lead.examiner@doi-state.test>\nSent: Monday, October 4, 2027\n"
                "To: Audit <audit@corp.test>\nSubject: FIN-801\n\nPlease send everything to me by 10/22/27.\n")
        who, _ = self.send_to(body, sender="ana.accountant@corp.test", to=(E.INTERN,), cc=())
        self.assertEqual(who, "lead.examiner@doi-state.test")

    def test_an_address_written_out(self):
        self.assertEqual(self.send_to("Please submit your documents to pat.kim@examfirm.test by Friday.")[0],
                         "pat.kim@examfirm.test")

    def test_default_is_the_lead_examiner(self):
        who, src = self.send_to("Please provide the claims listing by 10/20/27.")
        self.assertEqual(who, "lead.examiner@doi-state.test")
        self.assertIn("default", src)
        self.assertEqual(self.send_to("Please provide the claims listing.", leads=set())[0], "")   # none configured

    def test_name_not_on_the_email_is_flagged(self):
        from doi_analysis import analyze
        m = E.mail("FIN-801", "Please send the responses to Morgan by 10/20/27.", "lead.examiner@doi-state.test",
                   to=["audit@corp.test"])
        an = analyze(m, "FIN-801")
        self.assertEqual(an.send_to, "Morgan (address not on the email)")
        self.assertTrue(any("Morgan" in r for r in an.review))

    def test_not_a_send_instruction(self):
        for body in ("Please provide information related to Premium Tax by 10/20/27.",
                     "Please reconcile Schedule D to General Ledger balances.",
                     "Please provide responses to the following questions.",
                     "Please submit to SharePoint by Friday.",
                     "Please send to the Department by 10/20/27."):
            self.assertIn("default", self.send_to(body)[1], body)        # nothing stated -> default

    def test_a_colleagues_send_it_to_me_is_not_the_delivery_address(self):
        who, src = self.send_to("Please send it to me by Friday.", sender="ana.accountant@corp.test",
                                to=("pat.claims@corp.test",), cc=())
        self.assertIn("default", src)


class MustBeSentToDomainTests(unittest.TestCase):
    def test_a_domain_entry_means_anyone_at_that_domain(self):
        self.assertEqual(doi_filter._required("@Corp.test"), "@corp.test")
        self.assertEqual(doi_filter._required("corp.test"), "@corp.test")
        self.assertEqual(doi_filter._required(" Audit@Corp.test"), "audit@corp.test")
        with mock.patch.multiple(doi_filter, **dict(E.ROLES, MUST_BE_SENT_TO={"@corp.test"})):
            ok = detect(E.mail("FIN-901 Request", "Please provide the DOI examination support.",
                               "lead.examiner@doi-state.test", to=["pat.claims@corp.test"], cc=()))
            self.assertEqual(ok.status, "accept")
            no = detect(E.mail("FIN-901 Request", "Please provide the DOI examination support.",
                               "lead.examiner@doi-state.test", to=["someone@elsewhere.test"], cc=()))
            self.assertEqual(no.status, "reject")
            from upsert_tracker import contact
            m = E.mail("FIN-901", "Please provide the DOI items.", "ana.accountant@corp.test", to=["pat.claims@corp.test"], cc=())
            self.assertEqual(contact(m, "assignment"), "pat.claims@corp.test")   # a domain entry never hides colleagues


class SettingsLocationTests(unittest.TestCase):
    """local_settings.py next to the program, or in a Settings folder beside it."""

    def setUp(self):
        self.root = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.root)
        self.repo = self.root / "Repo"
        (self.root / "Settings").mkdir()
        self.repo.mkdir()

    def test_settings_folder_beside_the_program(self):
        beside = self.root / "Settings" / "local_settings.py"
        beside.write_text("X = 1\n")
        self.assertEqual(doi_filter.find_settings(self.repo), (beside, []))

    def test_program_folder_wins_and_the_other_copy_is_reported(self):
        here = self.repo / "local_settings.py"
        here.write_text("X = 1\n")
        beside = self.root / "Settings" / "local_settings.py"
        beside.write_text("X = 2\n")
        self.assertEqual(doi_filter.find_settings(self.repo), (here, [beside]))

    def test_none_found(self):
        self.assertEqual(doi_filter.find_settings(self.repo), (None, []))

    def test_settings_file_really_loads_from_the_settings_folder(self):
        import subprocess, sys
        src = Path(doi_filter.__file__).resolve().parent
        for f in src.glob("*.py"):
            if not f.name.startswith("local_settings"):
                shutil.copy(f, self.repo / f.name)
        (self.root / "Settings" / "local_settings.py").write_text(
            'TRACKER_PATH = "/fake/Tracker.xlsx"\nLOG_PATH = "/fake/CSV"\n_HELPER = 1\n')
        out = subprocess.run([sys.executable, "-c", "import doi_filter as F; print(F.TRACKER_PATH, F.LOG_PATH, hasattr(F, '_HELPER'))"],
                             cwd=self.repo, capture_output=True, text=True)
        self.assertEqual(out.stdout.split(), ["/fake/Tracker.xlsx", "/fake/CSV/DOI_Run_Log.csv", "False"])
        (self.root / "Settings" / "local_settings.py").unlink()
        out = subprocess.run([sys.executable, "run.py"], cwd=self.repo, capture_output=True, text=True)
        self.assertEqual(out.returncode, 2)
        self.assertIn("No local_settings.py found", out.stdout)


class SubmissionTests(unittest.TestCase):
    """'FIN-52 support has been uploaded to OneDrive' from inside the company to the examiner."""
    LEAD = "lead.examiner@doi-state.test"
    EXAMPLE = "Hello Lee,\n\nFIN-1052 support has been uploaded to OneDrive.\n\nHave a great weekend!\n"

    def setUp(self):
        p = mock.patch.multiple(doi_filter, **dict(E.ROLES, LEAD_EXAMINERS={self.LEAD}, MY_ADDRESSES={E.INTERN}))
        p.start()
        self.addCleanup(p.stop)

    def mail(self, body=EXAMPLE, sender="ana.accountant@corp.test", to=(LEAD,), subject="RE: FIN-1052 Request - Claims"):
        return E.mail(subject, body, sender, to=to, cc=(E.INTERN,), received=datetime(2027, 10, 8, 16, 0))

    def analyze(self, **kw):
        from doi_analysis import analyze
        return analyze(self.mail(**kw), "FIN-1052")

    def test_the_example_is_a_completed_submission(self):
        an = self.analyze()
        self.assertTrue(an.completes)
        self.assertIn("has been uploaded to OneDrive", an.submission_evidence)
        self.assertIn(self.LEAD, an.submitted_to)

    def test_the_example_is_accepted_without_doi_wording(self):
        det = detect(self.mail())
        self.assertEqual((det.status, det.request_id), ("accept", "FIN-1052"))

    def test_internal_handoff_is_not_a_submission(self):
        self.assertFalse(self.analyze(sender="pat.claims@corp.test", to=("ana.accountant@corp.test",)).completes)

    def test_examiner_uploading_is_not_a_submission(self):
        self.assertFalse(self.analyze(sender=self.LEAD, to=("ana.accountant@corp.test",)).completes)

    def test_partial_delivery(self):
        an = self.analyze(body="Hello Lee,\n\nPart of the FIN-1052 support has been uploaded; the remaining items will follow by Friday.\n")
        self.assertFalse(an.completes)
        self.assertTrue(an.partial_submission)

    def test_quoted_history_does_not_count(self):
        body = ("Thanks, will confirm Monday.\n\nFrom: Ana Accountant <ana.accountant@corp.test>\nSent: Friday\n"
                "To: Lee <lead.examiner@doi-state.test>\nSubject: RE: FIN-1052\n\nFIN-1052 support has been uploaded.\n")
        self.assertFalse(self.analyze(body=body).completes)


FIXTURE_PATH = Path(__import__("os").environ.get("DOI_TRACKER_FIXTURE", Path(__file__).parent / "fixtures" / "DOI_Exam_Test_Tracker_v2.xlsx"))


@unittest.skipUnless(FIXTURE_PATH.exists(), "needs a fake-data tracker (DOI_TRACKER_FIXTURE)")
class SubmissionTrackerTests(unittest.TestCase):
    """The tracker side: Status Completed, Completed On, never over a person's choice."""
    LEAD = SubmissionTests.LEAD

    def setUp(self):
        p = mock.patch.multiple(doi_filter, **dict(E.ROLES, LEAD_EXAMINERS={self.LEAD}, MY_ADDRESSES={E.INTERN}))
        p.start()
        self.addCleanup(p.stop)
        d = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, d)
        self.path = d / "Tracker.xlsx"
        shutil.copy(FIXTURE_PATH, self.path)

    def upsert(self, *mails):
        from upsert_tracker import upsert
        return upsert(self.path, [(m, "FIN-1052") for m in mails], backup_dir=self.path.parent / "Backups")

    def row(self, actions):
        from openpyxl import load_workbook
        ws = load_workbook(self.path)["DOI Exam Tracker"]
        r = actions[0]["row"]
        return ws, r

    def request(self):
        return E.mail("FIN-1052 Request - Claims", "Please provide the DOI examination claims listing by 10/20/27.",
                      self.LEAD, to=["ana.accountant@corp.test"], cc=[E.INTERN], eid="R1", received=datetime(2027, 10, 4, 9))

    def submission(self, body=SubmissionTests.EXAMPLE, eid="S1"):
        return E.mail("RE: FIN-1052 Request - Claims", body, "ana.accountant@corp.test", to=[self.LEAD],
                      cc=[E.INTERN], eid=eid, received=datetime(2027, 10, 8, 16))

    def test_submission_completes_the_row(self):
        self.upsert(self.request())
        acts = self.upsert(self.submission())
        ws, r = self.row(acts)
        self.assertEqual(ws[f"I{r}"].value, "Completed")
        self.assertEqual(ws[f"K{r}"].value, datetime(2027, 10, 8))
        self.assertIn("submitted to", ws[f"O{r}"].value)
        self.assertEqual(self.upsert(self.submission())[0]["action"], "unchanged")    # same email again: nothing

    def test_on_hold_is_left_alone(self):
        acts = self.upsert(self.request())
        ws, r = self.row(acts)
        ws[f"I{r}"] = "On Hold"
        ws.parent.save(self.path)
        ws, r = self.row(self.upsert(self.submission()))
        self.assertEqual(ws[f"I{r}"].value, "On Hold")
        self.assertIn("status left 'On Hold'", ws[f"O{r}"].value)

    def test_partial_submission_is_not_complete(self):
        self.upsert(self.request())
        ws, r = self.row(self.upsert(self.submission("Hello Lee,\n\nSome of the FIN-1052 support has been uploaded; the rest will follow.\n")))
        self.assertNotEqual(ws[f"I{r}"].value, "Completed")
        self.assertIsNone(ws[f"K{r}"].value)
        self.assertIn("partial submission", ws[f"O{r}"].value)
