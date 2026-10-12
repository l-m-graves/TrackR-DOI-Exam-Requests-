"""Tests for import_original.py. The 'original' list is built here from invented values that
mimic the real list's layout and quirks - no real data. Template tests need the fake tracker
fixture (see test_upsert_tracker.py)."""
import shutil
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from import_original import build, clean_request
from tests.test_upsert_tracker import FIXTURE
from upsert_tracker import SHEET, TABLE

HEADERS = ["Assigned to", "Initial Request Email Received", "Follow Up", "Request Name", "Due Date",
           "Contact", "Last Email Sent", "Status", "Completion Date"]
D = datetime


def fake_original(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Requests"
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    wb.create_sheet("Invoice")
    wb.save(path)


ROWS = [
    # assigned, received, follow up, request name, due, contact, last sent, status, completed
    ["Test Analyst 2", D(2027, 1, 5), None, "ABC-901 Request - Premium Tax Support", D(2027, 1, 20),
     "Alex Kim <alex.kim@example.test>", D(2027, 1, 18), "Completed", D(2027, 1, 19)],
    ["Test Analyst 2", D(2027, 1, 9), None, "ABC-902 Request - Reserve Detail - ABC-901 Follow-up", D(2027, 1, 30),
     "Alex Kim <alex.kim@example.test>; Sam Lee <sam.lee@example.test>", D(2027, 1, 25), "completed", D(2027, 1, 28)],
    ["Test Analyst 2", D(2027, 2, 1), None, "ABC-903: Reinsurance Schedules – ABC 902 Follow up", "TBD",
     "Jo Park", D(2027, 2, 3), "Sent to the team on 2/3; waiting on the reinsurance schedule.", None],
    ["Test Analyst 2", D(2027, 2, 2), None, "Quarterly Statement Questions", D(2027, 2, 20),
     None, None, "Completed", D(2027, 2, 10)],
]


class CleanRequestTests(unittest.TestCase):
    def test_tidy_name_unchanged(self):
        s = "ABC-902 Request - Reserve Detail - ABC-901 Follow-up"
        self.assertEqual(clean_request(s), (s, []))

    def test_colon_dash_spacing_and_follow_up_fixed(self):
        s, changes = clean_request("ABC-903: Reinsurance Schedules – ABC 902 Follow up")
        self.assertEqual(s, "ABC-903 Reinsurance Schedules - ABC-902 Follow-up")
        self.assertEqual(len(changes), 1)
        self.assertIn("tidied", changes[0])

    def test_missing_id_is_flagged(self):
        s, changes = clean_request("Quarterly Statement Questions")
        self.assertEqual(s, "Quarterly Statement Questions")
        self.assertTrue(any("no request ID" in c for c in changes))

    def test_blank(self):
        self.assertEqual(clean_request(None), ("", ["request name is blank"]))


@unittest.skipUnless(FIXTURE.exists(), f"fake tracker fixture not found at {FIXTURE}")
class BuildTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp)
        self.original = self.tmp / "original.xlsx"
        self.template = self.tmp / "template.xlsx"
        self.output = self.tmp / "out.xlsx"
        fake_original(self.original, ROWS)
        shutil.copy(FIXTURE, self.template)
        self.template_bytes = self.template.read_bytes()
        self.log, self.log_path = build(self.original, self.template, self.output)
        self.wb = load_workbook(self.output)
        self.ws = self.wb[SHEET]

    def test_inputs_untouched(self):
        self.assertEqual(self.template.read_bytes(), self.template_bytes)

    def test_table_holds_exactly_the_original_rows(self):
        self.assertEqual(self.ws.tables[TABLE].ref, "A3:O7")
        self.assertEqual(self.ws.tables[TABLE].autoFilter.ref, "A3:O7")
        self.assertEqual([self.ws[f"D{r}"].value.split()[0] for r in range(4, 8)],
                         ["ABC-901", "ABC-902", "ABC-903", "Quarterly"])
        for r in range(8, 15):                                     # template's example rows cleared
            self.assertTrue(all(self.ws.cell(r, c).value is None for c in range(1, 16)), r)
        self.assertFalse(any("DOI-TEST" in str(c.value) for row in self.ws.iter_rows() for c in row))

    def test_columns_mapped(self):
        ws = self.ws
        self.assertEqual(ws["B4"].value, "Test Analyst 2")
        self.assertEqual(ws["C4"].value, D(2027, 1, 5))
        self.assertEqual(ws["G4"].value, D(2027, 1, 20))
        self.assertEqual(ws["I4"].value, "Completed")
        self.assertEqual(ws["K4"].value, D(2027, 1, 19))
        self.assertEqual(ws["L4"].value, "Alex Kim <alex.kim@example.test>")
        self.assertEqual(ws["M4"].value, D(2027, 1, 18))
        self.assertIsNone(ws["O4"].value)
        self.assertEqual(ws["I5"].value, "Completed")              # 'completed' -> dropdown spelling
        template_fmt = load_workbook(self.template)[SHEET]["C4"].number_format
        self.assertEqual(ws["C4"].number_format, template_fmt)     # template formatting applied

    def test_formulas_on_every_row(self):
        for r in range(4, 8):
            self.assertEqual(self.ws[f"A{r}"].value, f'=IF($D{r}="","",LEFT($D{r},FIND(" ",$D{r}&" ")-1))')
            for col in "EFHJN":
                self.assertIn(f"$D{r}", self.ws[f"{col}{r}"].value)

    def test_free_text_status_and_text_due_date_go_to_notes(self):
        ws = self.ws
        self.assertIsNone(ws["I6"].value)
        self.assertIsNone(ws["G6"].value)
        self.assertIn("waiting on the reinsurance schedule", ws["O6"].value)
        self.assertIn("Due Date in original list: TBD", ws["O6"].value)

    def test_review_log(self):
        flagged = {x["tracker_row"]: x["review"] for x in self.log if x["review"]}
        self.assertEqual(sorted(flagged), [6, 7])
        self.assertIn("free text", flagged[6])
        self.assertIn("due date was text", flagged[6])
        self.assertIn("no request ID", flagged[7])
        self.assertTrue(self.log_path.exists())
        self.assertEqual([x["source_row"] for x in self.log], [2, 3, 4, 5])

    def test_new_assignee_added_to_dropdown_list(self):
        names = [c.value for c in self.wb["Lists"]["D"][3:60] if c.value]
        self.assertIn("Test Analyst 2", names)

    def test_summary_and_table_features_kept(self):
        self.assertEqual(self.wb["Summary"]["C11"].value, "=COUNTA(trkRequest)")
        tbl = self.ws.tables[TABLE]
        self.assertEqual(sum(c.calculatedColumnFormula is not None for c in tbl.tableColumns), 6)


if __name__ == "__main__":
    unittest.main()
