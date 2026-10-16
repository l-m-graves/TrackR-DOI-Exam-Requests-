"""Focused tests for doi_analysis.py (detection + analysis) and attachment_text.py."""
import shutil
import tempfile
import unittest
import zipfile
from datetime import date, datetime
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

import doi_filter
import tests.fake_doi_emails as E
from attachment_text import read_text
from doi_analysis import analyze, detect, segments
from tests.test_upsert_tracker import FIXTURE
from upsert_tracker import SHEET, upsert


class RoleSettingsMixin:
    def setUp(self):
        patcher = mock.patch.multiple(doi_filter, **E.ROLES)
        patcher.start()
        self.addCleanup(patcher.stop)


class DetectionTests(RoleSettingsMixin, unittest.TestCase):
    def test_direct_examiner_request(self):
        d = detect(E.DIRECT)
        self.assertEqual((d.status, d.confidence, d.request_id, d.sender_role), ("accept", "high", "FIN-201", "lead_examiner"))

    def test_forwarded_examiner_request_finds_the_original_sender(self):
        d = detect(E.FORWARDED)
        self.assertEqual((d.status, d.confidence), ("accept", "high"))
        self.assertEqual((d.sender_role, d.original_sender, d.original_role),
                         ("accountant", "examiner.two@doi-state.test", "examiner"))
        self.assertIn("forwarded", d.reasons[0])

    def test_cc_only_is_accepted(self):
        self.assertNotIn(E.INTERN, E.DIRECT.to)
        self.assertIn(E.INTERN, E.DIRECT.cc)
        self.assertEqual(detect(E.DIRECT).status, "accept")

    def test_internal_request_for_information(self):
        d = detect(E.INTERNAL_PAYROLL)
        self.assertEqual((d.status, d.confidence, d.request_id), ("accept", "medium", "FIN-203"))

    def test_request_id_only_in_an_attachment(self):
        self.assertEqual(detect(E.ATTACHMENT_ONLY).status, "review")        # before reading the file
        m = E.mail(E.ATTACHMENT_ONLY.subject, E.ATTACHMENT_ONLY.body, E.ATTACHMENT_ONLY.sender,
                   attachments=E.ATTACHMENT_ONLY.attachments, eid="E8b",
                   attachment_texts={"2027-10-04_DOI Request Letter.docx": E.ATTACHMENT_TEXT})
        d = detect(m)
        self.assertEqual((d.status, d.request_id), ("accept", "FIN-204"))
        self.assertIn("attachment", " ".join(d.reasons))

    def test_false_positives_are_rejected_with_a_reason(self):
        news = detect(E.FALSE_POSITIVE_NEWS)
        self.assertEqual(news.status, "reject")
        self.assertIn("no request ID", news.reasons[0])
        budget = detect(E.FALSE_POSITIVE_ID)
        self.assertEqual(budget.status, "reject")
        self.assertIn("nothing else says DOI exam", budget.reasons[0])

    def test_unconfigured_sender_goes_to_review_not_accept(self):
        d = detect(E.UNKNOWN_SENDER)
        self.assertEqual((d.status, d.request_id), ("review", "FIN-209"))
        self.assertIn("not in the role settings", d.reasons[0])

    def test_original_qualification_rules_still_accept_on_their_own(self):
        with mock.patch.object(doi_filter, "SENDER_DOMAINS", {"doi-state.test"}):
            d = detect(E.DIRECT)
        self.assertEqual((d.status, d.confidence), ("accept", "high"))
        self.assertIn("five qualification rules", d.reasons[0])


class AnalysisTests(RoleSettingsMixin, unittest.TestCase):
    def test_explicit_examiner_deadline(self):
        an = analyze(E.DIRECT, "FIN-201")
        self.assertEqual((an.exam_due_date, an.tracking_due_date, an.date_source, an.date_confidence),
                         (date(2027, 10, 22), date(2027, 10, 22), "examiner", "high"))
        self.assertIn("by October 22, 2027", an.date_evidence)
        self.assertEqual(an.external_contact, "lead.examiner@doi-state.test")

    def test_forwarded_deadline_is_the_examiners(self):
        an = analyze(E.FORWARDED, "FIN-202")
        self.assertEqual((an.exam_due_date, an.date_source), (date(2027, 10, 29), "examiner (forwarded)"))
        self.assertEqual((an.coordinator, an.external_contact, an.asked),
                         ("ana.accountant@corp.test", "examiner.two@doi-state.test", ["claims.lead@corp.test"]))
        self.assertEqual(an.departments, ["Payroll", "Claims"])

    def test_internal_expected_date_is_never_an_examiner_deadline(self):
        an = analyze(E.INTERNAL_PAYROLL, "FIN-203")
        self.assertIsNone(an.exam_due_date)
        self.assertEqual(an.internal_expected_date, date(2027, 10, 25))
        self.assertEqual((an.tracking_due_date, an.date_source, an.date_confidence),
                         (date(2027, 10, 25), "internal expected (fallback)", "low"))
        self.assertIn("not an examiner deadline", an.summary())
        self.assertEqual(an.departments, ["Payroll"])

    def test_conflicting_dates_are_kept_and_flagged(self):
        an = analyze(E.CONFLICTING, "FIN-207")
        self.assertIsNone(an.tracking_due_date)
        self.assertIn("conflicting examiner dates 2027-10-20", an.review[0])
        self.assertIn("2027-10-29", an.review[0])

    def test_details_only_in_attachment(self):
        m = E.mail(E.ATTACHMENT_ONLY.subject, E.ATTACHMENT_ONLY.body, E.ATTACHMENT_ONLY.sender,
                   attachments=E.ATTACHMENT_ONLY.attachments,
                   attachment_texts={"2027-10-04_DOI Request Letter.docx": E.ATTACHMENT_TEXT})
        an = analyze(m, "FIN-204")
        self.assertEqual((an.exam_due_date, an.date_source),
                         (date(2027, 11, 5), "attachment 2027-10-04_DOI Request Letter.docx"))
        self.assertEqual(an.request_items, ["Board meeting minutes for 2026", "The reinsurance contracts in force"])

    def test_multiple_items_and_scope(self):
        an = analyze(E.MULTI_ITEM, "FIN-208")
        self.assertEqual(len(an.request_items), 3)
        self.assertEqual(an.states, ["California", "Texas"])
        self.assertEqual(an.periods, ["January 1, 2026 to December 31, 2026"])
        self.assertEqual(an.controls, ["PT-07", "ITGC-04"])
        self.assertEqual((an.entities, an.products), (["Sample Casualty Company"], ["Personal Auto"]))

    def test_follow_up_links_to_the_earlier_request(self):
        self.assertEqual(analyze(E.FOLLOW_UP_SUBJECT, "FIN-205").parent_request_id, "FIN-201")
        an = analyze(E.FOLLOW_UP_BODY, "FIN-206")
        self.assertEqual((an.parent_request_id, an.request_type), ("FIN-201", "follow_up"))

    def test_other_department(self):
        an = analyze(E.IT_DEPARTMENT, "IT-39")
        self.assertEqual(an.departments, ["IT"])
        self.assertEqual(an.asked, ["it.lead@corp.test"])

    def test_clarification_missing_support_extension_meeting_submission(self):
        an = analyze(E.EVENTS, "FIN-201")
        for event in ("missing_support", "clarification", "extension", "meeting", "submission_instructions"):
            self.assertIn(event, an.events)
        self.assertEqual(an.exam_due_date, date(2027, 11, 12))            # the extended deadline
        self.assertEqual(analyze(E.FILES_READY, "FIN-201").submission_status, "files ready - not confirmed as submitted")
        self.assertEqual(analyze(E.SUBMITTED, "FIN-201").submission_status, "submitted (confirmed in email)")

    def test_segments_split_forward_headers(self):
        segs = segments(E.FORWARDED.body, E.FORWARDED.sender)
        self.assertEqual([s.author for s in segs], ["ana.accountant@corp.test", "examiner.two@doi-state.test"])
        self.assertTrue(segs[1].text.strip().startswith("Please provide the closed claims listing"))


class LastExchangeEmailTests(RoleSettingsMixin, unittest.TestCase):
    """Last Email Sent = the newest email in the request's exchange, whoever sent it."""

    def m(self, subject, sender, day, eid, conv=""):
        mail = E.mail(subject, "See attached.", sender, eid=eid, received=datetime(2027, 10, day, 15),
                      conversation_id=conv)
        mail.detection = detect(mail)
        return mail

    def test_newest_email_in_each_exchange_whoever_sent_it(self):
        request = self.m("FIN-201 Request - Premium Tax", "lead.examiner@doi-state.test", 4, "R1", conv="T201")
        thread = [
            request,
            self.m("RE: FIN-201 Request - Premium Tax", "ana.accountant@corp.test", 6, "A1", conv="T201"),
            self.m("Documents for the examiners", "pat.payroll@corp.test", 9, "P1", conv="T201"),   # subject changed
        ]
        others = [self.m("RE: IT-39 Request - Access Reviews", "examiner.two@doi-state.test", 7, "X1"),
                  self.m("Lunch on Friday", "ana.accountant@corp.test", 13, "A4", conv="T-other")]  # unrelated
        latest = doi_filter.last_exchange_email(thread + others, [(request, "FIN-201")])
        self.assertEqual(set(latest), {"FIN-201", "IT-39"})
        self.assertEqual(latest["FIN-201"], (datetime(2027, 10, 9, 15),
                                             "latest email in the exchange, from pat.payroll@corp.test"))
        self.assertEqual(latest["IT-39"], (datetime(2027, 10, 7, 15),
                                           "latest email in the exchange, from examiner.two@doi-state.test"))

    def test_examiner_reply_counts_too(self):
        request = self.m("FIN-201 Request - Premium Tax", "ana.accountant@corp.test", 4, "R1", conv="T201")
        reply = self.m("RE: FIN-201 Request - Premium Tax", "lead.examiner@doi-state.test", 11, "R2", conv="T201")
        self.assertEqual(doi_filter.last_exchange_email([request, reply], [(request, "FIN-201")])["FIN-201"][0],
                         datetime(2027, 10, 11, 15))

    def test_rejected_short_reply_still_counts(self):
        reply = self.m("RE: FIN-201 Request", "ana.accountant@corp.test", 8, "A5")
        reply.body = "Thanks."
        reply.detection = detect(reply)
        self.assertEqual(reply.detection.status, "reject")                 # no DOI wording - not a request
        self.assertIn("FIN-201", doi_filter.last_exchange_email([reply], []))

    def test_merge_keeps_newest_and_labels_sent_items(self):
        merged = doi_filter.merge_last_sent(
            {"FIN-201": (datetime(2027, 10, 9), "latest email in the exchange, from ana.accountant@corp.test")},
            {"FIN-201": datetime(2027, 10, 11), "IT-39": datetime(2027, 10, 2)})
        self.assertEqual(merged["FIN-201"], (datetime(2027, 10, 11), "your Sent Items"))
        self.assertEqual(merged["IT-39"], (datetime(2027, 10, 2), "your Sent Items"))
        self.assertEqual(doi_filter.merge_last_sent({"FIN-1": (datetime(2027, 1, 2), "a")}, None),
                         {"FIN-1": (datetime(2027, 1, 2), "a")})


class AttachmentTextTests(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.dir)

    def test_docx(self):
        p = self.dir / "letter.docx"
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("word/document.xml", "<w:document><w:body><w:p><w:r><w:t>Request FIN-204</w:t></w:r></w:p>"
                                            "<w:p><w:r><w:t>Due by November 5, 2027</w:t></w:r></w:p></w:body></w:document>")
        text, status = read_text(p)
        self.assertEqual(status, "read")
        self.assertEqual(text.split("\n")[:2], ["Request FIN-204", "Due by November 5, 2027"])

    def test_xlsx(self):
        p = self.dir / "items.xlsx"
        wb = Workbook(); wb.active.append(["1.", "Payroll register", "IT-39"]); wb.save(p)
        text, status = read_text(p)
        self.assertEqual((text, status), ("1. Payroll register IT-39", "read"))

    def test_pdf(self):
        p = self.dir / "request.pdf"
        content = b"BT /F1 12 Tf 72 720 Td (Request FIN-210 due 11/15/2027) Tj ET"
        objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
                b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
                b"/Resources << /Font << /F1 5 0 R >> >> >>",
                b"<< /Length %d >>stream\n" % len(content) + content + b"\nendstream",
                b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
        out, offsets = b"%PDF-1.4\n", []
        for i, o in enumerate(objs, 1):
            offsets.append(len(out)); out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
        xref = len(out)
        out += b"xref\n0 6\n0000000000 65535 f \n" + b"".join(b"%010d 00000 n \n" % o for o in offsets)
        out += b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % xref
        p.write_bytes(out)
        try:
            import pypdf  # noqa: F401
        except ImportError:
            self.skipTest("pypdf not installed")
        text, status = read_text(p)
        self.assertEqual(status, "read")
        self.assertIn("FIN-210", text)

    def test_unsupported_missing_and_corrupt_files_are_reported(self):
        (self.dir / "old.doc").write_bytes(b"\xd0\xcf")
        (self.dir / "shot.png").write_bytes(b"\x89PNG")
        (self.dir / "bad.docx").write_bytes(b"not a zip")
        self.assertTrue(read_text(self.dir / "old.doc")[1].startswith("not read - .doc files are not supported"))
        self.assertTrue(read_text(self.dir / "shot.png")[1].startswith("not read - .png"))
        self.assertTrue(read_text(self.dir / "bad.docx")[1].startswith("not read -"))
        with mock.patch.dict("sys.modules", {"pypdf": None}):
            self.assertIn("pip install pypdf", read_text(self.dir / "x.pdf")[1])


@unittest.skipUnless(FIXTURE.exists(), f"fake tracker fixture not found at {FIXTURE}")
class WorkbookTests(RoleSettingsMixin, unittest.TestCase):
    def setUp(self):
        super().setUp()
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp)
        self.path = self.tmp / "tracker.xlsx"
        shutil.copy(FIXTURE, self.path)

    def run_upsert(self, *pairs):
        for m, rid in pairs:
            m.analysis = analyze(m, rid)
        return upsert(self.path, list(pairs), backup_dir=self.tmp / "Backups")

    def ws(self):
        return load_workbook(self.path)[SHEET]

    def test_forwarded_request_row(self):
        [a] = self.run_upsert((E.FORWARDED, "FIN-202"))
        ws, r = self.ws(), a["row"]
        self.assertEqual(ws[f"G{r}"].value, datetime(2027, 10, 29))
        self.assertEqual(ws[f"L{r}"].value, "claims.lead@corp.test")          # the person being asked
        self.assertIn("exam due 2027-10-29 (examiner (forwarded))", ws[f"O{r}"].value)

    def test_internal_fallback_is_labelled_in_notes(self):
        [a] = self.run_upsert((E.INTERNAL_PAYROLL, "FIN-203"))
        notes = self.ws()[f"O{a['row']}"].value
        self.assertEqual(self.ws()[f"G{a['row']}"].value, datetime(2027, 10, 25))
        self.assertIn("Due Date set from the internal expected date - no examiner deadline seen yet", notes)

    def test_follow_up_row_feeds_follows_up_on(self):
        [a] = self.run_upsert((E.FOLLOW_UP_BODY, "FIN-206"))
        self.assertEqual(self.ws()[f"D{a['row']}"].value, "FIN-206 Request - Additional Support - FIN-201 Follow-up")

    def test_last_email_sent_is_written_with_its_source(self):
        [a] = self.run_upsert((E.DIRECT, "FIN-201"))
        acts = upsert(self.path, [], backup_dir=self.tmp / "Backups", last_sent={
            "FIN-201": (datetime(2027, 10, 9, 15, 30), "latest email in the exchange, from lead.examiner@doi-state.test")})
        self.assertEqual(acts[0]["changes"],
                         "last email sent 2027-10-09 (latest email in the exchange, from lead.examiner@doi-state.test)")
        self.assertEqual(self.ws()[f"M{a['row']}"].value, datetime(2027, 10, 9))

    def test_new_examiner_date_on_a_dated_row_is_flagged_not_chosen(self):
        [a] = self.run_upsert((E.DIRECT, "FIN-201"))
        [b] = self.run_upsert((E.EVENTS, "FIN-201"))                         # extension to 11/12/2027
        ws = self.ws()
        self.assertEqual(ws[f"G{a['row']}"].value, datetime(2027, 10, 22))   # kept
        self.assertIn("email states due 2027-11-12; tracker has 2027-10-22", b["review"])
        self.assertIn("REVIEW:", ws[f"O{a['row']}"].value)


if __name__ == "__main__":
    unittest.main()
