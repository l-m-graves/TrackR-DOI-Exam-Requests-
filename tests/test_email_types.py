"""Tests for email_types.py and the template-driven tracker updates."""
import shutil
import tempfile
import unittest
from datetime import datetime
from unittest import mock
from pathlib import Path

from openpyxl import load_workbook

import doi_filter
from doi_filter import Mail, SentMail
from due_dates import find_due, find_expected
from email_types import classify, topic_from_subject
from tests.fake_templates import (ACKNOWLEDGMENT, ASSIGNMENT, RESPONSE_MATERIALS, SCOPE_QUESTION,
                                  STATUS_UPDATE)
from tests.test_upsert_tracker import FIXTURE
from upsert_tracker import SHEET, request_text, upsert

DAY = datetime(2027, 10, 4, 9, 30)


class ClassifyTests(unittest.TestCase):
    def test_each_template_is_recognised(self):
        for (subject, body), kind in [(ASSIGNMENT, "assignment"), (ACKNOWLEDGMENT, "acknowledgment"),
                                      (STATUS_UPDATE, "status_update"), (RESPONSE_MATERIALS, "response_materials"),
                                      (SCOPE_QUESTION, "scope_question")]:
            with self.subTest(kind=kind):
                self.assertEqual(classify(subject, body), kind)

    def test_reply_quoting_the_assignment_is_not_an_assignment(self):
        self.assertEqual(classify("RE: DOI Exam Request FIN-990 - Topic",
                                  "Thanks, noted.\n\nFrom: Sam\nSent: Monday\nThe DOI has issued examination request FIN-990."),
                         "other")

    def test_topic_and_column_d_text(self):
        self.assertEqual(topic_from_subject(ASSIGNMENT[0], "FIN-990"), "Premium Tax Support")
        self.assertEqual(topic_from_subject("RE: " + ASSIGNMENT[0], "FIN-990"), "Premium Tax Support")
        self.assertIsNone(topic_from_subject(STATUS_UPDATE[0], "FIN-990"))
        self.assertIsNone(topic_from_subject(SCOPE_QUESTION[0], "FIN-990"))
        self.assertEqual(request_text(ASSIGNMENT[0], "FIN-990"), "FIN-990 Request - Premium Tax Support")

    def test_dates_in_each_template(self):
        self.assertEqual(str(find_due(ASSIGNMENT[1], DAY)["date"]), "2027-10-22")
        self.assertEqual(str(find_due(ACKNOWLEDGMENT[1], DAY)["date"]), "2027-10-22")
        self.assertEqual(str(find_expected(STATUS_UPDATE[1], DAY)), "2027-10-29")


@unittest.skipUnless(FIXTURE.exists(), f"fake tracker fixture not found at {FIXTURE}")
class TemplateFlowTests(unittest.TestCase):
    """One request's life: assignment received -> acknowledgment, status update, question and
    response materials sent."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp)
        self.path = self.tmp / "tracker.xlsx"
        shutil.copy(FIXTURE, self.path)
        # The templates' team is internal here (example.test is also the default DOI test domain).
        p = mock.patch.multiple(doi_filter, SENDER_DOMAINS={"doi-state.test"}, INTERNAL_DOMAINS={"example.test"})
        p.start()
        self.addCleanup(p.stop)

    def assignment(self):
        subject, body = ASSIGNMENT
        return Mail(subject, body, "sam.coordinator@example.test", ["jo.park@example.test", "lee.kim@example.test"],
                    ["manager@example.test"], ["FIN-990 Request.pdf"], DAY, "A990", "Sam Coordinator",
                    names={"jo.park@example.test": "Jo Park"}, conversation_id="T990")

    def sent(self, template, day, eid):
        return (SentMail(template[0], datetime(2027, 10, day, 15), "T990", eid, template[1]), "FIN-990")

    def run_upsert(self, accepted=(), sent=()):
        return upsert(self.path, list(accepted), backup_dir=self.tmp / "Backups", sent=list(sent))

    def ws(self):
        return load_workbook(self.path)[SHEET]

    def test_assignment_creates_a_complete_row(self):
        [a] = self.run_upsert([(self.assignment(), "FIN-990")])
        ws, r = self.ws(), a["row"]
        self.assertEqual(a["action"], "added")
        self.assertEqual(ws[f"D{r}"].value, "FIN-990 Request - Premium Tax Support")
        self.assertEqual(ws[f"G{r}"].value, datetime(2027, 10, 22))
        self.assertEqual(ws[f"I{r}"].value, "Not Started")
        self.assertEqual(ws[f"L{r}"].value, "Jo Park <jo.park@example.test>; lee.kim@example.test")   # assignees
        self.assertTrue(ws[f"O{r}"].value.startswith("2027-10-04 Assignment from sam.coordinator@example.test"))

    def test_sent_templates_progress_the_row(self):
        [a] = self.run_upsert([(self.assignment(), "FIN-990")])
        r = a["row"]
        acts = self.run_upsert(sent=[self.sent(ACKNOWLEDGMENT, 5, "S1"), self.sent(STATUS_UPDATE, 12, "S2"),
                                     self.sent(SCOPE_QUESTION, 13, "S3"), self.sent(RESPONSE_MATERIALS, 20, "S4")])
        self.assertEqual(len(acts), 4)
        self.assertIn("status In Progress (acknowledgment)", acts[0]["changes"])
        self.assertNotIn("status In Progress", acts[1]["changes"])       # already In Progress
        ws = self.ws()
        self.assertEqual(ws[f"I{r}"].value, "In Progress")
        self.assertEqual(ws[f"G{r}"].value, datetime(2027, 10, 22))    # status update's date did NOT replace it
        notes = ws[f"O{r}"].value.split("\n")
        self.assertEqual(len(notes), 5)
        self.assertTrue(notes[1].startswith("2027-10-05 Acknowledgment sent:"))
        self.assertIn("expects to provide by 2027-10-29", notes[2])
        self.assertTrue(notes[3].startswith("2027-10-13 Scope question sent:"))
        self.assertTrue(notes[4].startswith("2027-10-20 Response materials sent:"))
        # running again changes nothing
        self.assertEqual(self.run_upsert(sent=[self.sent(ACKNOWLEDGMENT, 5, "S1"), self.sent(STATUS_UPDATE, 12, "S2")]), [])

    def test_status_chosen_by_a_person_is_never_changed(self):
        [a] = self.run_upsert([(self.assignment(), "FIN-990")])
        wb = load_workbook(self.path)
        wb[SHEET][f"I{a['row']}"] = "On Hold"
        wb.save(self.path)
        [s] = self.run_upsert(sent=[self.sent(STATUS_UPDATE, 12, "S2")])
        self.assertNotIn("status In Progress", s["changes"])
        self.assertEqual(self.ws()[f"I{a['row']}"].value, "On Hold")

    def test_follow_up_email_for_unknown_request_does_not_create_a_row(self):
        subject, body = STATUS_UPDATE
        m = Mail(subject.replace("FIN-990", "FIN-991"), body, "jo.park@example.test", [], [], ["x.pdf"], DAY, "U1")
        before = self.ws().tables["tblRequests"].ref
        [a] = self.run_upsert([(m, "FIN-991")])
        self.assertEqual(a["action"], "skipped")
        self.assertIn("not in the tracker", a["changes"])
        self.assertEqual(self.ws().tables["tblRequests"].ref, before)


if __name__ == "__main__":
    unittest.main()
