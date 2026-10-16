"""Tests for upsert_tracker.py against a copy of a FAKE-DATA tracker.

The workbook itself is not in the repo (*.xlsx is git-ignored). Put a copy that holds
only test data at tests/fixtures/DOI_Exam_Test_Tracker_v2.xlsx, or point the
DOI_TRACKER_FIXTURE environment variable at one. Without it these tests are skipped.
Each test works on a temporary copy - the fixture is never modified.
"""
import os
import shutil
import tempfile
import unittest
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator

from doi_filter import REQUEST_ID, Mail
from upsert_tracker import (SHEET, TABLE, TrackerError, first_word_id, note_ref, request_text, upsert)

FIXTURE = Path(os.environ.get("DOI_TRACKER_FIXTURE",
                              Path(__file__).parent / "fixtures" / "DOI_Exam_Test_Tracker_v2.xlsx"))


def mail(subject, received, entry_id, sender="examiner@example.test", name="Jordan Lee", body="please provide"):
    return Mail(subject=subject, body=body, sender=sender, to=["audit@example.test"],
                attachments=["Request.pdf"], received=received, entry_id=entry_id, sender_name=name)


class PureHelperTests(unittest.TestCase):
    def test_request_text_keeps_clean_subject(self):
        s = "DOI-TEST-60 Request - Claims - TEST-41 Follow-up"
        self.assertEqual(request_text(s, "DOI-TEST-60"), s)

    def test_request_text_strips_reply_prefixes(self):
        self.assertEqual(request_text("RE: FW: DOI-TEST-60 Request", "DOI-TEST-60"), "DOI-TEST-60 Request")

    def test_request_text_puts_id_first(self):
        self.assertEqual(request_text("Docs for DOI-TEST-60", "DOI-TEST-60"), "DOI-TEST-60 Docs for DOI-TEST-60")
        self.assertEqual(request_text("DOI-TEST-60 - docs", "DOI-TEST-60"), "DOI-TEST-60 - docs")
        self.assertEqual(request_text("DOI-TEST-60: docs", "DOI-TEST-60"), "DOI-TEST-60: docs")  # column A strips ':' 

    def test_first_word_id_matches_column_a_formula(self):
        self.assertEqual(first_word_id("FIN-52: Fake Subject"), "FIN-52")      # ':' stripped
        self.assertEqual(first_word_id("IT-39 Request"), "IT-39")
        self.assertEqual(first_word_id("Quarterly Statement Questions"), "")   # sheet shows (no ID)
        self.assertEqual(first_word_id(None), "")

    def test_request_text_keeps_colon_style(self):
        self.assertEqual(request_text("RE: FIN-52: Fake Subject", "FIN-52"), "FIN-52: Fake Subject")

    def test_note_ref_is_stable_and_short(self):
        self.assertEqual(note_ref("ABC"), note_ref("ABC"))
        self.assertNotEqual(note_ref("ABC"), note_ref("ABD"))
        self.assertEqual(len(note_ref("ABC")), len("ref ") + 8)


@unittest.skipUnless(FIXTURE.exists(), f"fake tracker fixture not found at {FIXTURE}")
class UpsertTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp)
        self.path = self.tmp / "tracker.xlsx"
        shutil.copy(FIXTURE, self.path)
        self.backups = self.tmp / "backups"

    def run_upsert(self, mails, **kw):
        accepted = [(m, REQUEST_ID.search(m.subject).group(0)) for m in mails]
        return upsert(self.path, accepted, backup_dir=self.backups, **kw)

    def sheet(self):
        return load_workbook(self.path)[SHEET]

    def table_ref(self):
        return self.sheet().tables[TABLE].ref

    def row_of_or_none(self, rid):
        try:
            return self.row_of(rid)
        except StopIteration:
            return None

    def row_of(self, rid):
        ws = self.sheet()
        return next(r for r in range(4, ws.max_row + 1) if str(ws[f"D{r}"].value or "").startswith(rid + " "))

    # --- new rows ---------------------------------------------------------------
    def test_new_id_appends_row_with_formulas_and_extends_table(self):
        before = self.table_ref()
        m = mail("DOI-TEST-60 Request - Claims Handling", datetime(2027, 10, 8, 9, 30), "FAKE60")
        [a] = self.run_upsert([m])
        self.assertEqual(a["action"], "added")
        ws = self.sheet()
        r = a["row"]
        self.assertEqual(before, f"A3:O{r - 1}")
        self.assertEqual(ws.tables[TABLE].ref, f"A3:O{r}")
        self.assertEqual(ws.tables[TABLE].autoFilter.ref, f"A3:O{r}")
        # formulas copied down from the first data row and shifted to the new row
        self.assertEqual(ws[f"A{r}"].value, Translator(ws["A4"].value, origin="A4").translate_formula(f"A{r}"))
        for col in "EFHJN":
            self.assertTrue(str(ws[f"{col}{r}"].value).startswith("="), col)
            self.assertIn(f"$D{r}", ws[f"{col}{r}"].value)
            self.assertNotIn(f"$D{r - 1}", ws[f"{col}{r}"].value)
        # inputs
        self.assertEqual(ws[f"C{r}"].value, datetime(2027, 10, 8))           # date only, like the sheet
        self.assertEqual(ws[f"C{r}"].number_format, ws["C4"].number_format)  # style copied from first row
        self.assertEqual(ws[f"D{r}"].value, "DOI-TEST-60 Request - Claims Handling")
        self.assertEqual(ws[f"I{r}"].value, "Not Started")
        self.assertIsNone(ws[f"L{r}"].value)      # the examiner is never the Contact; nobody internal was asked yet
        self.assertIn(note_ref("FAKE60"), ws[f"O{r}"].value)
        # columns left for a person
        for col in "BGKM":
            self.assertIsNone(ws[f"{col}{r}"].value, col)

    def test_two_new_ids_in_one_run(self):
        acts = self.run_upsert([mail("DOI-TEST-60 Request", datetime(2027, 10, 8), "A"),
                                mail("DOI-TEST-61 Request", datetime(2027, 10, 9), "B")])
        self.assertEqual([a["row"] for a in acts], [acts[0]["row"], acts[0]["row"] + 1])
        self.assertEqual(self.table_ref(), f"A3:O{acts[1]['row']}")

    def test_reply_for_new_id_in_same_batch_updates_not_duplicates(self):
        acts = self.run_upsert([mail("RE: DOI-TEST-60 Request", datetime(2027, 10, 9), "B"),
                                mail("DOI-TEST-60 Request", datetime(2027, 10, 8), "A")])
        self.assertEqual([a["action"] for a in acts], ["added", "updated"])   # processed oldest first
        ws = self.sheet()
        r = acts[0]["row"]
        self.assertEqual(ws[f"D{r}"].value, "DOI-TEST-60 Request")
        self.assertEqual(ws[f"O{r}"].value.count("ref "), 2)

    # --- existing rows ----------------------------------------------------------
    def test_existing_id_appends_note_and_keeps_everything_else(self):
        ws0 = self.sheet()
        n = self.row_of("DOI-TEST-50")
        ref_before = self.table_ref()
        old = {c: ws0[f"{c}{n}"].value for c in "BCDGILMO"}
        m = mail("RE: DOI-TEST-50 Request - Technology Risk Review", datetime(2027, 10, 9, 16, 0), "FAKE50",
                 sender="someone.else@example.test", name="Someone Else")
        [a] = self.run_upsert([m])
        self.assertEqual((a["action"], a["row"], a["changes"]), ("updated", n, "note"))
        ws = self.sheet()
        for c in "BCDGILM":
            self.assertEqual(ws[f"{c}{n}"].value, old[c], c)                  # untouched
        self.assertTrue(ws[f"O{n}"].value.startswith(old["O"] + "\n"))         # appended, not overwritten
        self.assertIn("[" + note_ref("FAKE50") + "]", ws[f"O{n}"].value)
        self.assertEqual(self.table_ref(), ref_before)                         # no new row

    def test_reply_matches_row_whose_id_has_a_colon(self):
        [a] = self.run_upsert([mail("FIN-9901: Fake Reserve Detail", datetime(2027, 10, 8), "C1")])
        self.assertEqual(self.sheet()[f"D{a['row']}"].value, "FIN-9901: Fake Reserve Detail")
        [b] = self.run_upsert([mail("RE: FIN-9901 Fake Reserve Detail", datetime(2027, 10, 9), "C2")])
        self.assertEqual((b["action"], b["row"]), ("updated", a["row"]))     # same row, no duplicate

    # --- due date and last email sent --------------------------------------------
    def test_new_row_gets_due_date_from_email(self):
        m = mail("DOI-TEST-60 Request", datetime(2027, 10, 8), "D1", body="Please provide by October 22, 2027.")
        [a] = self.run_upsert([m])
        self.assertIn("due date", a["changes"])
        self.assertEqual(self.sheet()[f"G{a['row']}"].value, datetime(2027, 10, 22))

    def test_existing_due_date_is_never_overwritten_only_noted(self):
        n = self.row_of("DOI-TEST-50")
        before = self.sheet()[f"G{n}"].value
        self.assertIsNotNone(before)
        [a] = self.run_upsert([mail("RE: DOI-TEST-50 Request", datetime(2027, 10, 9), "D2",
                                    body="Extension granted - now due 11/30/2027.")])
        ws = self.sheet()
        self.assertEqual(ws[f"G{n}"].value, before)
        self.assertNotIn("due date", a["changes"])
        self.assertIn("email states due 2027-11-30; tracker has", ws[f"O{n}"].value)

    def test_asap_goes_to_notes_not_the_date_column(self):
        [a] = self.run_upsert([mail("DOI-TEST-61 Request", datetime(2027, 10, 8), "D3", body="Please respond ASAP.")])
        ws = self.sheet()
        self.assertIsNone(ws[f"G{a['row']}"].value)
        self.assertIn("due: ASAP", ws[f"O{a['row']}"].value)

    def test_last_email_sent_only_moves_forward(self):
        n = self.row_of("DOI-TEST-50")
        cur = self.sheet()[f"M{n}"].value
        older, newer = cur.replace(day=1) if cur.day > 1 else cur, datetime(2028, 1, 15, 14, 5)
        self.assertEqual(self.run_upsert([], last_sent={"DOI-TEST-50": older}), [])     # not older
        [a] = self.run_upsert([], last_sent={"DOI-TEST-50": newer, "DOI-TEST-999": newer})
        self.assertEqual((a["request_id"], a["changes"], a["mail"]), ("DOI-TEST-50", "last email sent 2028-01-15", None))
        self.assertEqual(self.sheet()[f"M{n}"].value, datetime(2028, 1, 15))           # date only
        self.assertIsNone(self.row_of_or_none("DOI-TEST-999"))                         # no row created
        self.assertEqual(self.run_upsert([], last_sent={"DOI-TEST-50": newer}), [])    # idempotent

    def test_earlier_email_moves_initial_received_back(self):
        [a] = self.run_upsert([mail("DOI-TEST-50 Request", datetime(2027, 10, 1, 8), "EARLY")])
        self.assertIn("received date", a["changes"])
        self.assertEqual(self.sheet()[f"C{self.row_of('DOI-TEST-50')}"].value, datetime(2027, 10, 1))

    # --- idempotency, backups, safety ------------------------------------------
    def test_second_run_changes_nothing_and_makes_no_backup(self):
        mails = [mail("DOI-TEST-60 Request", datetime(2027, 10, 8), "A"),
                 mail("RE: DOI-TEST-50 Request", datetime(2027, 10, 9), "B")]
        self.run_upsert(mails)
        first_bytes = self.path.read_bytes()
        self.assertEqual(len(list(self.backups.iterdir())), 1)
        acts = self.run_upsert(mails)
        self.assertEqual({a["action"] for a in acts}, {"unchanged"})
        self.assertEqual(self.path.read_bytes(), first_bytes)               # file not even re-saved
        self.assertEqual(len(list(self.backups.iterdir())), 1)

    def test_backup_is_the_pre_write_file(self):
        original = self.path.read_bytes()
        self.run_upsert([mail("DOI-TEST-60 Request", datetime(2027, 10, 8), "A")])
        [b] = self.backups.iterdir()
        self.assertEqual(b.read_bytes(), original)

    def test_dry_run_writes_nothing(self):
        original = self.path.read_bytes()
        [a] = self.run_upsert([mail("DOI-TEST-60 Request", datetime(2027, 10, 8), "A")], dry_run=True)
        self.assertEqual(a["action"], "added")
        self.assertEqual(self.path.read_bytes(), original)
        self.assertFalse(self.backups.exists())

    def test_refuses_when_excel_has_file_open(self):
        (self.tmp / "~$tracker.xlsx").write_bytes(b"")
        with self.assertRaises(TrackerError):
            self.run_upsert([mail("DOI-TEST-60 Request", datetime(2027, 10, 8), "A")])

    def test_refuses_workbook_with_chart(self):
        with zipfile.ZipFile(self.path, "a") as z:
            z.writestr("xl/charts/chart1.xml", "<chartSpace/>")
        with self.assertRaises(TrackerError):
            self.run_upsert([mail("DOI-TEST-60 Request", datetime(2027, 10, 8), "A")])

    def test_workbook_features_survive_save(self):
        self.run_upsert([mail("DOI-TEST-60 Request", datetime(2027, 10, 8), "A")])
        wb = load_workbook(self.path)
        ws = wb[SHEET]
        tbl = ws.tables[TABLE]
        self.assertEqual(tbl.tableStyleInfo.name, "TableStyleMedium2")
        self.assertEqual(sum(c.calculatedColumnFormula is not None for c in tbl.tableColumns), 6)
        self.assertIn("StatusList", wb.defined_names)
        self.assertIn("trkRequest", wb.defined_names)
        self.assertTrue(any(str(dv.sqref) == "I4:I1000" for dv in ws.data_validations.dataValidation))
        self.assertGreater(len(list(ws.conditional_formatting)), 0)
        self.assertIsNotNone(ws["A3"].comment)
        self.assertEqual(wb["Summary"]["C11"].value, "=COUNTA(trkRequest)")


if __name__ == "__main__":
    unittest.main()
